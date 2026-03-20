import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


PROJECT_ROOT = Path(__file__).resolve().parent
OUTPUT_PATH = PROJECT_ROOT / "class_descriptions.docx"
MODEL_ROOTS = [
    PROJECT_ROOT / "src" / "Backend" / "Services" / "UserService" / "Models",
    PROJECT_ROOT
    / "src"
    / "Backend"
    / "Services"
    / "ListingService"
    / "ListingService.Domain"
    / "Models",
]
BASE_TYPE_ROOTS = [
    PROJECT_ROOT
    / "src"
    / "Backend"
    / "Services"
    / "ListingService"
    / "ListingService.Domain"
    / "Abstractions"
]
MAX_FIELDS_PER_CLASS = 8

SCALAR_TYPES = {
    "string",
    "bool",
    "int",
    "long",
    "short",
    "byte",
    "decimal",
    "double",
    "float",
    "Guid",
    "DateTime",
    "DateTimeOffset",
}
IGNORED_FIELD_NAMES = {"ListingId", "ListingDetailsType"}
IGNORED_EXACT_CLASS_NAMES = {"OfficeClassId"}
ALLOWED_FOREIGN_KEY_FIELDS = {"Id", "AgentId"}
FIELD_PRIORITY = {
    "Id": 100,
    "Name": 95,
    "UserName": 95,
    "Email": 95,
    "Title": 95,
    "Description": 90,
    "AddressText": 90,
    "AgentId": 88,
    "Rooms": 88,
    "Floor": 84,
    "Floors": 84,
    "PlotArea": 84,
    "HouseArea": 84,
    "PriceAmount": 88,
    "PriceCurrency": 84,
    "PricePeriod": 82,
    "AreaTotal": 86,
    "PublishedDate": 84,
    "Status": 82,
    "BuildYear": 80,
    "Created": 78,
    "ExpiresAt": 78,
    "ChangedAt": 78,
    "AccountImageUrl": 75,
    "GoogleId": 75,
    "TokenHash": 75,
}
DOMAIN_CLASS_HINTS = {
    "User",
    "RefreshToken",
    "GoogleAuthentication",
    "Listing",
    "ApartmentDetails",
    "HouseDetails",
    "LandDetails",
    "OfficeDetails",
    "WarehouseDetails",
    "ListingMedia",
    "ListingPriceHistory",
}


@dataclass
class PropertyInfo:
    name: str
    type_name: str
    declared_in: str


@dataclass
class ClassInfo:
    name: str
    file_path: Path
    base_types: list[str] = field(default_factory=list)
    interfaces: list[str] = field(default_factory=list)
    properties: list[PropertyInfo] = field(default_factory=list)
    inferred_properties: list[PropertyInfo] = field(default_factory=list)


def normalize_type_name(type_name: str) -> str:
    clean_name = type_name.replace("?", "").strip()
    clean_name = re.sub(r"\s+", "", clean_name)
    return clean_name


def is_scalar_type(type_name: str) -> bool:
    clean_name = normalize_type_name(type_name)
    if clean_name in SCALAR_TYPES:
        return True
    return clean_name.startswith("Nullable<") and clean_name[9:-1] in SCALAR_TYPES


def split_base_types(raw_bases: str | None) -> tuple[list[str], list[str]]:
    if not raw_bases:
        return [], []

    bases: list[str] = []
    interfaces: list[str] = []
    for part in raw_bases.split(","):
        clean = part.strip()
        clean = re.sub(r"<.*?>", "", clean)
        if not clean:
            continue
        if clean.startswith("I"):
            interfaces.append(clean)
        else:
            bases.append(clean)
    return bases, interfaces


def parse_class_file(file_path: Path) -> ClassInfo | None:
    text = file_path.read_text(encoding="utf-8")
    class_match = re.search(
        r"public\s+(?:sealed\s+)?class\s+(?P<name>\w+)\s*(?::\s*(?P<bases>[^{]+))?",
        text,
    )
    if not class_match:
        return None

    class_name = class_match.group("name")
    raw_bases = class_match.group("bases")
    base_types, interfaces = split_base_types(raw_bases)

    property_pattern = re.compile(
        r"public\s+(?:required\s+)?(?P<type>[\w<>\[\]\?,\.]+)\s+"
        r"(?P<name>\w+)\s*\{\s*get;\s*set;\s*\}",
        re.MULTILINE,
    )
    properties = [
        PropertyInfo(
            name=match.group("name"),
            type_name=match.group("type"),
            declared_in=class_name,
        )
        for match in property_pattern.finditer(text)
    ]

    inferred_properties = infer_object_initializer_properties(text, class_name)
    return ClassInfo(
        name=class_name,
        file_path=file_path,
        base_types=base_types,
        interfaces=interfaces,
        properties=properties,
        inferred_properties=inferred_properties,
    )


def infer_object_initializer_properties(text: str, class_name: str) -> list[PropertyInfo]:
    inferred: list[PropertyInfo] = []
    pattern = re.compile(
        rf"new\s+{class_name}\s*\{{(?P<body>.*?)\}}",
        re.DOTALL,
    )
    match = pattern.search(text)
    if not match:
        return inferred

    body = match.group("body")
    for assignment in re.finditer(r"(?P<name>\w+)\s*=", body):
        property_name = assignment.group("name")
        if property_name in {item.name for item in inferred}:
            continue
        inferred.append(
            PropertyInfo(
                name=property_name,
                type_name="inferred",
                declared_in=class_name,
            )
        )
    return inferred


def load_classes() -> dict[str, ClassInfo]:
    classes: dict[str, ClassInfo] = {}
    for root in [*MODEL_ROOTS, *BASE_TYPE_ROOTS]:
        if not root.exists():
            continue
        for file_path in root.rglob("*.cs"):
            class_info = parse_class_file(file_path)
            if class_info:
                classes[class_info.name] = class_info
    return classes


def collect_inherited_properties(
    class_info: ClassInfo,
    class_map: dict[str, ClassInfo],
    visited: set[str] | None = None,
) -> list[PropertyInfo]:
    if visited is None:
        visited = set()

    inherited: list[PropertyInfo] = []
    for base_name in class_info.base_types:
        if base_name in visited:
            continue
        visited.add(base_name)
        base_info = class_map.get(base_name)
        if not base_info:
            continue
        inherited.extend(base_info.properties)
        inherited.extend(collect_inherited_properties(base_info, class_map, visited))
    return inherited


def is_meaningful_property(property_info: PropertyInfo) -> bool:
    if property_info.name in IGNORED_FIELD_NAMES:
        return False

    if property_info.name == "Id":
        return True

    if property_info.name.endswith("By") and property_info.name in {"CreatedBy", "LastModifiedBy"}:
        return False

    clean_type = normalize_type_name(property_info.type_name)
    if (
        property_info.name.endswith("Id")
        and property_info.name not in ALLOWED_FOREIGN_KEY_FIELDS
        and clean_type in {"Guid", "Nullable<Guid>"}
    ):
        return False

    if property_info.type_name == "inferred":
        return True

    if clean_type.startswith("ICollection<") or clean_type.startswith("List<"):
        return False

    return is_scalar_type(clean_type)


def property_score(property_info: PropertyInfo) -> int:
    score = FIELD_PRIORITY.get(property_info.name, 50)
    if property_info.name.endswith("Id") and property_info.name != "Id":
        score += 8
    if property_info.name.startswith("Has"):
        score += 4
    if property_info.name in {"Created", "LastModified", "ListingDetailsType"}:
        score -= 10
    return score


def class_score(class_info: ClassInfo) -> int:
    score = 0
    scalar_count = sum(1 for item in class_info.properties if is_meaningful_property(item))
    score += scalar_count * 3

    if class_info.name in DOMAIN_CLASS_HINTS:
        score += 12
    if class_info.name == "Listing":
        score += 14
    if class_info.name == "User":
        score += 14
    if class_info.name.endswith("Details"):
        score += 10
    if "ILookupEntity" in class_info.interfaces:
        score -= 20
    if class_info.name.endswith(("Type", "Category", "Material", "Class")):
        score -= 12
    if class_info.name in IGNORED_EXACT_CLASS_NAMES:
        score -= 20
    return score


def select_main_classes(class_map: dict[str, ClassInfo]) -> list[ClassInfo]:
    candidates = [
        class_info
        for class_info in class_map.values()
        if class_info.file_path.is_relative_to(MODEL_ROOTS[0])
        or class_info.file_path.is_relative_to(MODEL_ROOTS[1])
    ]
    ranked = sorted(
        candidates,
        key=lambda item: (-class_score(item), item.name),
    )

    selected: list[ClassInfo] = []
    for class_info in ranked:
        if class_score(class_info) < 10:
            continue
        selected.append(class_info)

    preferred_order = [
        "User",
        "RefreshToken",
        "GoogleAuthentication",
        "Listing",
        "ApartmentDetails",
        "HouseDetails",
        "LandDetails",
        "OfficeDetails",
        "WarehouseDetails",
        "ListingMedia",
        "ListingPriceHistory",
    ]
    order_map = {name: index for index, name in enumerate(preferred_order)}
    selected.sort(key=lambda item: (order_map.get(item.name, 999), item.name))
    return selected


def deduplicate_properties(properties: Iterable[PropertyInfo]) -> list[PropertyInfo]:
    unique: list[PropertyInfo] = []
    seen: set[str] = set()
    for property_info in properties:
        if property_info.name in seen:
            continue
        seen.add(property_info.name)
        unique.append(property_info)
    return unique


def extract_main_properties(class_info: ClassInfo, class_map: dict[str, ClassInfo]) -> list[PropertyInfo]:
    inherited_properties = collect_inherited_properties(class_info, class_map)
    all_properties = deduplicate_properties(
        [*class_info.properties, *inherited_properties, *class_info.inferred_properties]
    )
    meaningful = [item for item in all_properties if is_meaningful_property(item)]
    order_map = {item.name: index for index, item in enumerate(all_properties)}
    sorted_by_priority = sorted(
        meaningful,
        key=lambda prop: (-property_score(prop), order_map[prop.name]),
    )

    if len(sorted_by_priority) <= MAX_FIELDS_PER_CLASS:
        return sorted_by_priority

    return sorted_by_priority[:MAX_FIELDS_PER_CLASS]


def split_identifier(name: str) -> list[str]:
    return re.findall(r"[A-Z]?[a-z]+|[A-Z]+(?=[A-Z]|$)|\d+", name)


def description_from_name(field_name: str) -> str:
    exact_descriptions = {
        "Id": "уникальный идентификатор записи",
        "UserName": "логин пользователя для идентификации в системе",
        "Name": "наименование или отображаемое имя сущности",
        "Email": "адрес электронной почты пользователя",
        "AccountImageUrl": "ссылка на изображение профиля пользователя",
        "TokenHash": "хешированное значение токена для безопасной проверки",
        "ExpiresAt": "дата и время окончания срока действия токена",
        "GoogleId": "идентификатор пользователя во внешнем сервисе Google",
        "AgentId": "идентификатор агента, связанного с объявлением",
        "AddressText": "текстовое представление адреса объекта",
        "Title": "краткий заголовок объявления",
        "Description": "подробное описание объекта или записи",
        "PriceAmount": "числовое значение цены объекта",
        "PriceCurrency": "валюта, в которой указана цена",
        "PricePeriod": "периодичность применения указанной цены",
        "AreaTotal": "общая площадь объекта",
        "BuildYear": "год постройки объекта",
        "PublishedDate": "дата и время публикации объявления",
        "Status": "текущее состояние или статус записи",
        "Rooms": "общее количество комнат",
        "Bedrooms": "количество спален",
        "FloorsTotal": "общее количество этажей в здании",
        "Levels": "количество уровней помещения",
        "Floor": "этаж расположения объекта",
        "AreaLiving": "жилая площадь объекта",
        "AreaKitchen": "площадь кухни",
        "HasBalcony": "признак наличия балкона",
        "HasLoggia": "признак наличия лоджии",
        "HasElevator": "признак наличия лифта",
        "CeilingHeight": "высота потолков объекта",
        "Floors": "количество этажей объекта",
        "HouseArea": "площадь дома",
        "PlotArea": "площадь земельного участка",
        "HasElectricity": "признак наличия электроснабжения",
        "HasGarage": "признак наличия гаража",
        "HasBasement": "признак наличия подвала",
        "HasFence": "признак наличия ограждения участка",
        "AllowedUse": "разрешённый вид использования земельного участка",
        "HasGas": "признак наличия газоснабжения",
        "HasWater": "признак наличия водоснабжения",
        "HasSewer": "признак наличия канализации",
        "HasBuildings": "признак наличия построек на участке",
        "ParkingSpaces": "количество доступных парковочных мест",
        "Kitchenette": "признак наличия мини-кухни",
        "FloorLoadTonsPerSquareMeter": "допустимая нагрузка на пол в тоннах на квадратный метр",
        "GateCount": "количество ворот на объекте",
        "HasDockLevelers": "признак наличия доковых уравнительных платформ",
        "DockCount": "количество доковых зон",
        "HasRamps": "признак наличия пандусов",
        "PowerKw": "доступная электрическая мощность в киловаттах",
        "HasOffice": "признак наличия офисной зоны",
        "OfficeArea": "площадь офисной части объекта",
        "IsMainMedia": "признак основного медиафайла объявления",
        "Url": "адрес размещения медиафайла",
        "ChangedAt": "дата и время изменения цены",
        "Created": "дата и время создания записи",
    }
    if field_name in exact_descriptions:
        return exact_descriptions[field_name]

    if field_name.endswith("Id") and field_name != "Id":
        target_name = split_identifier(field_name[:-2])
        target_text = " ".join(part.lower() for part in target_name)
        return f"идентификатор связанной сущности {target_text}"

    if field_name.startswith("Has"):
        feature = " ".join(part.lower() for part in split_identifier(field_name[3:]))
        return f"признак наличия характеристики {feature}"

    words = " ".join(part.lower() for part in split_identifier(field_name))
    return f"техническое поле для хранения значения {words}"


def configure_base_styles(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)


def add_class_heading(document: Document, class_name: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(class_name)
    run.bold = True
    run.font.size = Pt(13)


def add_field_paragraph(document: Document, index: int, field_name: str, description: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.add_run(f"{index} {field_name} – {description}.")


def create_word_document(class_sections: list[tuple[str, list[PropertyInfo]]]) -> None:
    document = Document()
    configure_base_styles(document)

    for class_name, properties in class_sections:
        if not properties:
            continue
        add_class_heading(document, class_name)
        for index, property_info in enumerate(properties, start=1):
            description = description_from_name(property_info.name)
            add_field_paragraph(document, index, property_info.name, description)

    document.save(OUTPUT_PATH)


def build_document_sections(class_map: dict[str, ClassInfo]) -> list[tuple[str, list[PropertyInfo]]]:
    sections: list[tuple[str, list[PropertyInfo]]] = []
    for class_info in select_main_classes(class_map):
        properties = extract_main_properties(class_info, class_map)
        if properties:
            sections.append((class_info.name, properties))
    return sections


def main() -> None:
    class_map = load_classes()
    sections = build_document_sections(class_map)
    create_word_document(sections)
    print(f"Файл сохранён: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
